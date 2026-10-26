import frappe
import base64
import io
import re
from typing import List, Dict, Optional, Any
import os
import hashlib
try:
    from PIL import Image
except ImportError:
    pass

class AttachmentEvidence:
    def __init__(self, url: str, doc_type: str, status: str = "EXTRACTED", raw_text: str = "", file_hash: Optional[str] = None):
        self.url = url
        self.document_type = doc_type
        self.status = status
        self.raw_text = raw_text
        self.fields = {}
        self.file_hash = file_hash
        
    def add_field(self, field_name: str, value: Any, confidence: float = 1.0, status: str = "EXTRACTED", reason: str = None):
        self.fields[field_name] = {
            "field": field_name,
            "value": value,
            "confidence": confidence,
            "source": self.url,
            "status": status,
            "reason": reason
        }

    def to_dict(self):
        data = {
            "document_type": self.document_type,
            "extraction_status": self.status,
            "source_attachment": self.url,
            "extracted_text_content": self.raw_text,
            "extracted_fields": list(self.fields.values())
        }
        if self.file_hash:
            data["file_hash"] = self.file_hash
        return data

TravelReportEvidence = AttachmentEvidence

class UniversalAttachmentProcessor:
    @staticmethod
    def classify_document(url: str, text: str = "", explicit_type: str = "") -> str:
        name = url.lower()
        content = text.lower()
        combined = f"{name} {content} {explicit_type.lower()}"
        
        if "travel authorization" in combined or "authorization" in name:
            return "TRAVEL_AUTHORIZATION"
        if "vendor travel log" in combined or ("log" in combined and "vendor" in combined) or "vehicle log" in combined or "log book" in combined:
            return "VENDOR_TRAVEL_LOG"
        if "travel report" in name or "travel_report" in name or "travel report" in content or "odometer" in content:
            return "TRAVEL_REPORT"
        if "invoice" in combined or "bill" in name or "receipt" in combined:
            # Need to separate fuel/toll from generic vendor invoice
            if "toll" in combined: return "TOLL_RECEIPT"
            if "fuel" in combined: return "FUEL_RECEIPT"
            return "VENDOR_INVOICE"
        if "boarding" in combined or "pass" in name:
            return "BOARDING_PASS"
        if "ticket" in combined or "flight" in combined or "air" in name:
            return "TICKET"
        if "trip report" in combined or "training report" in combined or "visit report" in combined or "monitoring plan" in combined:
            return "TRIP_REPORT"
        if "gps" in combined or "location" in combined:
            return "GPS_EVIDENCE"
        return "OTHER"

    @staticmethod
    def extract_pdf_text(content: bytes) -> str:
        text = ""
        try:
            import PyPDF2
            import io
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            for page in reader.pages:
                text += page.extract_text() + "\n"
        except Exception as e:
            frappe.log_error(f"Error extracting PDF text: {e}")
        return text

    @staticmethod
    def extract_docx_text(content: bytes) -> str:
        text = ""
        try:
            import docx
            import io
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([para.text for para in doc.paragraphs])
            for table in doc.tables:
                for row in table.rows:
                    text += " | ".join([cell.text for cell in row.cells]) + "\n"
        except Exception as e:
            frappe.log_error(f"Error extracting DOCX text: {e}")
        return text

    @staticmethod
    def extract_msg_text(content: bytes) -> str:
        text = ""
        try:
            import extract_msg
            import io
            msg = extract_msg.Message(io.BytesIO(content))
            text = f"Subject: {msg.subject}\nDate: {msg.date}\nBody: {msg.body}"
        except Exception as e:
            frappe.log_error(f"Error extracting MSG text: {e}")
        return text

    @staticmethod
    def extract_embedded_images_from_pdf(content: bytes, max_images: int = 6) -> List[Dict]:
        """
        Extracts embedded raster images (e.g., odometer photos, vehicle photos) from a PDF.
        Filters out small icons, signatures, logos (< 150x150).
        """
        extracted = []
        try:
            import fitz
            doc = fitz.open(stream=content, filetype="pdf")
            for page_idx, page in enumerate(doc):
                image_list = page.get_images(full=True)
                for img_idx, img_info in enumerate(image_list):
                    if len(extracted) >= max_images:
                        break
                    xref = img_info[0]
                    base_img = doc.extract_image(xref)
                    width = base_img.get("width", 0)
                    height = base_img.get("height", 0)
                    if width < 150 or height < 150:
                        continue
                    
                    try:
                        pix = fitz.Pixmap(doc, xref)
                        if pix.n >= 5:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                        
                        max_dim = max(pix.width, pix.height)
                        if max_dim > 900:
                            try:
                                from PIL import Image
                                img = Image.open(io.BytesIO(base_img["image"]))
                                if img.mode != "RGB":
                                    img = img.convert("RGB")
                                img.thumbnail((900, 900), Image.Resampling.LANCZOS)
                                buf = io.BytesIO()
                                img.save(buf, format="JPEG", quality=75)
                                b64_str = base64.b64encode(buf.getvalue()).decode("utf-8")
                            except Exception:
                                jpeg_bytes = pix.tobytes("jpeg")
                                b64_str = base64.b64encode(jpeg_bytes).decode("utf-8")
                        else:
                            jpeg_bytes = pix.tobytes("jpeg")
                            b64_str = base64.b64encode(jpeg_bytes).decode("utf-8")
                            
                        extracted.append({
                            "page": page_idx + 1,
                            "image_index": img_idx + 1,
                            "image_data": b64_str,
                            "source_type": "embedded_pdf_image"
                        })
                    except Exception as e:
                        frappe.log_error(f"Error processing PDF image xref {xref}: {e}")
            doc.close()
        except Exception as e:
            frappe.log_error(f"Error extracting embedded PDF images: {e}")
        return extracted

    @staticmethod
    def extract_embedded_images_from_docx(content: bytes, max_images: int = 6) -> List[Dict]:
        """
        Extracts embedded images from docx media directory (word/media/).
        """
        extracted = []
        try:
            import zipfile
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                media_files = [n for n in z.namelist() if n.startswith("word/media/") and n.lower().endswith((".png", ".jpg", ".jpeg", ".webp"))]
                for idx, media_name in enumerate(media_files):
                    if len(extracted) >= max_images:
                        break
                    img_bytes = z.read(media_name)
                    try:
                        from PIL import Image
                        img = Image.open(io.BytesIO(img_bytes))
                        if img.width < 150 or img.height < 150:
                            continue
                        if img.mode != "RGB":
                            img = img.convert("RGB")
                        img.thumbnail((900, 900), Image.Resampling.LANCZOS)
                        buf = io.BytesIO()
                        img.save(buf, format="JPEG", quality=75)
                        b64_str = base64.b64encode(buf.getvalue()).decode("utf-8")
                    except Exception:
                        b64_str = base64.b64encode(img_bytes).decode("utf-8")

                    extracted.append({
                        "page": 1,
                        "image_index": idx + 1,
                        "image_data": b64_str,
                        "source_type": "embedded_docx_image"
                    })
        except Exception as e:
            frappe.log_error(f"Error extracting embedded DOCX images: {e}")
        return extracted

    @staticmethod
    def _extract_travel_report_fields(text: str, evidence: AttachmentEvidence):
        text_lower = text.lower()
        
        start_match = re.search(r'start(?:ing)?\s+odometer[:\s]*(\d+)', text_lower)
        if start_match:
            val = float(start_match.group(1))
            evidence.add_field("start_odometer_text", val, 0.95, "VERIFIED", evidence.url)
            evidence.add_field("start_odometer", val, 0.95, "VERIFIED", evidence.url)
        else:
            evidence.add_field("start_odometer", None, 0.0, "NOT_FOUND", evidence.url)
            
        end_match = re.search(r'end(?:ing)?\s+odometer[:\s]*(\d+)', text_lower)
        if end_match:
            val = float(end_match.group(1))
            evidence.add_field("end_odometer_text", val, 0.95, "VERIFIED", evidence.url)
            evidence.add_field("end_odometer", val, 0.95, "VERIFIED", evidence.url)
        else:
            evidence.add_field("end_odometer", None, 0.0, "NOT_FOUND", evidence.url)
            
        km_match = re.search(r'(?:total\s+)?km[:\s]*(\d+)', text_lower)
        if km_match: evidence.add_field("claimed_km", float(km_match.group(1)), 0.90, "EXTRACTED", evidence.url)
            
        date_match = re.search(r'date[:\s]*([\d\-/]+)', text_lower)
        if date_match: evidence.add_field("travel_date", date_match.group(1), 0.90, "EXTRACTED", evidence.url)
        
        if "gps" in text_lower or "location" in text_lower:
            evidence.add_field("gps", True, 0.9, "EXTRACTED", evidence.url)

    @staticmethod
    def _extract_vendor_invoice_fields(text: str, evidence: AttachmentEvidence):
        text_lower = text.lower()
        amt_match = re.search(r'(?:total|amount|grand total|sum)[\s:a-z]*([\d,]+\.\d{2})', text_lower)
        if amt_match: evidence.add_field("amount", float(amt_match.group(1).replace(",", "")), 0.8, "EXTRACTED", evidence.url)
        else: evidence.add_field("amount", None, 0.0, "NOT_FOUND", evidence.url)
        
        date_match = re.search(r'date[:\s]*([\d\-/]+)', text_lower)
        if date_match: evidence.add_field("invoice_date", date_match.group(1), 0.9, "EXTRACTED", evidence.url)

    @staticmethod
    def process_attachments(file_dicts: List[Any], max_pages_per_pdf: Optional[int] = None) -> Dict[str, Any]:
        images_for_vision = []
        structured_evidence = []
        
        for fdict in file_dicts:
            url = fdict.get("url") if isinstance(fdict, dict) else fdict
            explicit_type = fdict.get("explicit_type", "") if isinstance(fdict, dict) else ""
            if not url: continue
            
            try:
                import requests
                import io
                base_url = "https://erp.micromerger.com"
                fetch_url = base_url + url if url.startswith("/") else url
                
                content = None
                # Check local disk first
                try:
                    file_docs = frappe.get_all("File", filters={"file_url": url}, fields=["name"], limit=1)
                    if file_docs:
                        fdoc = frappe.get_doc("File", file_docs[0].name)
                        fpath = fdoc.get_full_path()
                        if os.path.exists(fpath):
                            with open(fpath, "rb") as f:
                                content = f.read()
                except Exception:
                    pass

                if not content:
                    try:
                        response = requests.get(fetch_url, timeout=10)
                        response.raise_for_status()
                        content = response.content
                    except Exception as e:
                        frappe.log_error(f"Failed to fetch remote file {fetch_url}: {e}", "Attachment Processor")
                        continue
                
                file_hash = hashlib.sha256(content).hexdigest() if content else None
                ext = url.lower()
                text_content = ""
                has_usable_text = False
                embedded_images = []
                
                if ext.endswith(".pdf"):
                    text_content = UniversalAttachmentProcessor.extract_pdf_text(content)
                    text_content = text_content[:5000]
                    clean_text = text_content.strip().lower()
                    if len(clean_text) > 50 and clean_text.count("camscanner") < (len(clean_text) / 20):
                        has_usable_text = True
                    embedded_images = UniversalAttachmentProcessor.extract_embedded_images_from_pdf(content)
                    
                elif ext.endswith(".docx"):
                    text_content = UniversalAttachmentProcessor.extract_docx_text(content)
                    text_content = text_content[:5000]
                    has_usable_text = len(text_content.strip()) > 50
                    embedded_images = UniversalAttachmentProcessor.extract_embedded_images_from_docx(content)
                    
                elif ext.endswith(".msg"):
                    text_content = UniversalAttachmentProcessor.extract_msg_text(content)
                    text_content = text_content[:5000]
                    has_usable_text = len(text_content.strip()) > 50
                
                doc_type = UniversalAttachmentProcessor.classify_document(url, text_content, explicit_type)
                evidence = AttachmentEvidence(url, doc_type, raw_text=text_content, file_hash=file_hash)
                
                if has_usable_text:
                    if doc_type == "TRAVEL_REPORT":
                        UniversalAttachmentProcessor._extract_travel_report_fields(text_content, evidence)
                    elif doc_type == "VENDOR_INVOICE":
                        UniversalAttachmentProcessor._extract_vendor_invoice_fields(text_content, evidence)
                
                structured_evidence.append(evidence.to_dict())
                
                # Add any embedded images from the report to vision queue
                if embedded_images:
                    for img_item in embedded_images:
                        images_for_vision.append({
                            "url": url,
                            "page": img_item.get("page", 1),
                            "image_index": img_item.get("image_index", 1),
                            "image_data": img_item["image_data"],
                            "doc_type": doc_type,
                            "source_type": img_item.get("source_type", "embedded_report_image")
                        })
                
                needs_vision = ext.endswith((".png", ".jpg", ".jpeg", ".webp")) or (ext.endswith(".pdf") and not has_usable_text and not embedded_images)
                if doc_type in ["VENDOR_INVOICE", "TOLL_RECEIPT", "FUEL_RECEIPT", "TICKET", "VENDOR_TRAVEL_LOG"]:
                    needs_vision = True
                    
                if needs_vision:
                    if ext.endswith(".pdf"):
                        from ai_analytics.new_api import ImageProcessor
                        max_pages = max_pages_per_pdf if max_pages_per_pdf is not None else 3
                        rendered = ImageProcessor.process_pdf_pages(content, max_pages=max_pages)
                        for page_idx, b64_str in enumerate(rendered):
                            images_for_vision.append({
                                "url": url,
                                "page": page_idx + 1,
                                "image_data": b64_str,
                                "doc_type": doc_type,
                                "source_type": "rendered_pdf_page"
                            })
                    elif ext.endswith((".png", ".jpg", ".jpeg", ".webp")):
                        try:
                            from PIL import Image
                            img = Image.open(io.BytesIO(content))
                            if img.mode != "RGB": img = img.convert("RGB")
                            img.thumbnail((900, 900), Image.Resampling.LANCZOS)
                            buffer = io.BytesIO()
                            img.save(buffer, format="JPEG", quality=75)
                            b64_str = base64.b64encode(buffer.getvalue()).decode("utf-8")
                        except Exception:
                            b64_str = base64.b64encode(content).decode("utf-8")
                            
                        images_for_vision.append({
                            "url": url,
                            "page": 1,
                            "image_data": b64_str,
                            "doc_type": doc_type,
                            "source_type": "standalone_image"
                        })
                        
            except Exception as e:
                frappe.log_error(f"Failed to process attachment {url}: {e}")
                
        return {
            "images_for_vision": images_for_vision,
            "structured_evidence": structured_evidence
        }
